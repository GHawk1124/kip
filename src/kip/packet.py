"""One component packet convention: ordered work, known sheets, visible gaps."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from .content import Column, Sources, Table
from .doc.blocks import Block
from .doc.graph import analyze_code
from .doc.loader import KipSyntaxError
from .doc.validate import Diagnostic
from .req import Item, Requirement, Requirements, requirements_table, variables_table, compliance_matrix
from .sheets import Constants, Sheet
from .packet_config import PacketConfig, SHEETS, STAGES


class UnavailableInput(Exception):
    """A declared input is absent; it is not a Python or schema error."""


class MissingInput:
    def __init__(self, reason):
        self.reason = reason

    def __getattr__(self, name):
        raise UnavailableInput(self.reason)

    def __getitem__(self, key):
        raise UnavailableInput(self.reason)


@dataclass
class Material:
    stage: str
    value: object = None
    state: str = "OPEN"
    reason: str = ""
    optional: bool = False


def _sheet(root, spec):
    """Validate headers even in a blank starter; empty valid sheets are OPEN."""
    from openpyxl import load_workbook
    from .authoring import read_records

    path = root / spec.path
    if not path.exists():
        raise UnavailableInput(f"Missing {spec.path} ({spec.sheet}).")
    book = load_workbook(path, read_only=True)
    try:
        if spec.sheet not in book.sheetnames:
            raise UnavailableInput(f"Missing sheet {spec.sheet!r} in {spec.path}.")
        headers = next(book[spec.sheet].values, ())
        if len(headers) != len(set(headers)) or any(not isinstance(h, str) or not h for h in headers):
            raise ValueError(f"{spec.path}/{spec.sheet}: headers must be unique, nonempty text")
        missing = set(spec.columns) - set(headers)
        if missing:
            raise ValueError(f"{spec.path}/{spec.sheet}: missing columns {', '.join(sorted(missing))}")
    finally:
        book.close()
    records = read_records(path, spec.sheet)
    table = Sheet(records, {h: h for h in headers}, path)
    key = next((k for k in ("key", "id", "number") if k in headers), None)
    if key:
        table.require_unique(key)
    for field in spec.required:
        table.require_present(field)
    if not records:
        raise UnavailableInput(f"{spec.path}/{spec.sheet} has no rows yet.")
    return table


class ComponentPacket:
    def __init__(self, blocks, path):
        declarations = [b for b in blocks if b.kind == "packet"]
        marker = self.marker = declarations[0]
        self.path = Path(path)
        self.root = self.path.resolve().parent
        self.authored = [b for b in blocks if b.kind != "packet"]
        self.materials = {}
        self.status = {}
        if len(declarations) != 1 or marker.id != "component" or marker.source.strip():
            self.error('use one empty "# %% packet component" declaration')
        if set(marker.meta) - {"na", "config"}:
            self.error("packet accepts na= and config=; put report settings in run_document")
        config_path = self.root / marker.meta.get("config", "packet.toml")
        try:
            if "config" in marker.meta and not config_path.exists():
                raise ValueError("configuration file does not exist")
            self.config = PacketConfig(config_path)
        except (ValueError, OSError, TypeError) as exc:
            self.error(f"{config_path}: {exc}")
        self.sections = self.config.sections
        self.sheets = self.config.sheets
        self.na = self.config.na | set(filter(None, (s.strip() for s in marker.meta.get("na", "").split(","))))
        if self.na - self.sections.keys():
            self.error("na= sections must appear in the packet order")
        reserved = {alias for alias, name in (("C", "constants"), ("reqs", "requirements")) if name in self.sheets}
        for block in self.authored:
            if block.id.startswith("_packet_"):
                self.error("block ids beginning _packet_ are reserved")
            if analyze_code(block.source)[0] & reserved and block.is_code:
                self.error("component supplies C and reqs; remove their duplicate assignments")
            if block.kind == "prelude":
                continue
            stage = block.meta.setdefault("stage", self.config.defaults[block.kind])
            if stage not in self.sections:
                self.error(f"block {block.id}: unknown or removed stage {stage!r}; set stage= or change packet defaults")
            if stage in self.na:
                self.error(f"{stage} is N/A but contains authored block {block.id}")

    def error(self, message):
        raise KipSyntaxError(message, self.path, self.marker.marker_line)

    def load(self, doc):
        self.materials = {}
        self.identity = None
        # Rebuilds re-read inputs and discard earlier packet diagnostics.
        doc.diagnostics[:] = [d for d in doc.diagnostics if d.block_id != self.marker.id]
        for name, spec in self.sheets.items():
            stage = spec.stage
            material = Material(stage, optional=spec.optional)
            try:
                value = self.load_requirements() if name == "requirements" else _sheet(self.root, spec)
                if stage in self.na:
                    raise ValueError(f"{stage} is N/A but {name} contains data")
                if name == "constants":
                    value = Constants.load(self.root / spec.path, spec.sheet)
                elif name == "references":
                    value = Sources(value.sources())
                material.value, material.state = value, "PRESENT"
                if name == "tests" and any(not str(r.get(k, "")).strip() for r in value for k in ("result", "evidence")):
                    material.state, material.reason = "OPEN", "Test results or evidence are not recorded for every row."
            except UnavailableInput as exc:
                material.reason = str(exc)
            except Exception as exc:
                material.state, material.reason = "ERROR", f"{type(exc).__name__}: {exc}"
                doc.diagnostics.append(Diagnostic("error", self.marker.id, self.marker.marker_line, material.reason))
            self.materials[name] = material
        reqs = self.value("requirements")
        tests = self.value("tests")
        if reqs is not None and tests is not None:
            unknown = {str(r["requirement"]) for r in tests} - set(reqs.all_requirements())
            if unknown:
                reason = self.sheets["tests"].path + ": unknown requirement(s) " + ", ".join(sorted(unknown))
                self.materials["tests"].state, self.materials["tests"].reason = "ERROR", reason
                doc.diagnostics.append(Diagnostic("error", self.marker.id, self.marker.marker_line, reason))
        return {alias: self.materials[name].value if self.materials[name].value is not None
                else MissingInput(self.materials[name].reason)
                for alias, name in (("C", "constants"), ("reqs", "requirements")) if name in self.materials}

    def value(self, name):
        material = self.materials.get(name)
        return material.value if material else None

    def load_requirements(self):
        spec = self.sheets["requirements"]
        toml = self.root / self.config.requirements_file
        xlsx = self.root / spec.path
        if toml.exists() and xlsx.exists():
            raise ValueError(f"choose {self.config.requirements_file} or {spec.path}; both are present")
        if toml.exists():
            reqs = Requirements.load(toml)
        elif xlsx.exists():
            sheet = _sheet(self.root, spec)
            reqs = Requirements(Item(self.root.name), {
                str(r["id"]): Requirement(str(r["id"]), str(r["text"]), str(r["verification"])) for r in sheet
            }, {}, path=xlsx)
        else:
            raise UnavailableInput(f"Missing {self.config.requirements_file} or {spec.path}.")
        self.identity = reqs.item
        for variable in reqs.all_variables().values():
            variable.quantity  # Validate even when no calculation uses the variable.
        if any(not r.text.strip() for r in reqs.requirements.values()):
            raise ValueError("requirement text must not be blank")
        if not reqs.requirements:
            raise UnavailableInput("No local requirements are declared yet.")
        return reqs

    def finish(self, doc):
        """Assemble views after execution, so compliance uses final checks."""
        from .doc.kernel import BlockResult
        from .prose import literal
        from .quantities import symbol_table, controlled_entry

        stages = {s: [] for s in self.sections}
        for block in self.authored:
            if block.kind != "prelude":
                stages[block.meta["stage"]].append(block)
        self.status = {}
        pending = set()
        for stage in self.sections:
            work = [doc.results[b.id] for b in stages[stage]]
            required = [m for m in self.materials.values() if m.stage == stage and not m.optional]
            states = [m.state for m in required] + [r.state if r.ok else "ERROR" for r in work]
            if any(s != "PRESENT" for s in states):
                pending.add(stage)
            present = any(r.latex or r.typst or r.text or r.content or r.checks for r in work)
            state = next((s for s in ("ERROR", "BLOCKED", "OPEN") if s in states), "PRESENT" if required or present else "OPEN")
            reason = "; ".join(m.reason for m in [*required, *work] if m.reason)
            if state == "BLOCKED":
                reason = "Dependent work is waiting on missing inputs."
            if not required and not present and not reason:
                reason = "No work supplied yet."
            self.status[stage] = ("N/A", "Explicitly excluded in packet settings.") if stage in self.na else (state, reason)
        if "overview" in self.sections and self.sections["overview"].generate and "overview" not in self.na:
            if "overview" not in pending:
                self.status["overview"] = ("PRESENT", "") if self.identity else ("OPEN", "Item identity is not declared in requirements.")
        reqs = self.value("requirements")
        if "compliance" in self.sections and "compliance" not in self.na:
            status = self.compliance_status(doc, reqs, stages["compliance"])
            if "compliance" not in pending or status[0] in ("FAIL", "BLOCKED"):
                self.status["compliance"] = status

        entries = [item for b in self.authored for item in doc.results[b.id].quantities.values()]
        constants = self.value("constants")
        if constants is not None:
            entries = [constants.entry(k) for k in constants] + entries
        if reqs:
            entries += [controlled_entry(v) for v in reqs.all_variables().values()]
        symbols = None
        if entries and "inputs" in self.sections and self.sections["inputs"].generate and "inputs" not in self.na:
            try:
                symbols = symbol_table(entries)
            except ValueError as exc:
                self.status["inputs"] = ("ERROR", str(exc))
                doc.diagnostics.append(Diagnostic("error", self.marker.id, self.marker.marker_line, str(exc)))
        rendered = [b for b in self.authored if b.kind == "prelude"]

        def add(stage, name, content=None, *, state="PRESENT", reason="", kind="table", heading=False):
            bid = "_packet_" + name
            title = self.sections[stage].title if heading else name.removeprefix("material_").replace("_", " ").title()
            meta = {"label": title, "wrap_title": "false"}
            if name in ("identity", "overview") or name.endswith("_status"):
                meta.pop("label")
            if heading:
                meta.update(section="1", columns=str(self.sections[stage].columns))
            block = Block(bid, "text" if heading else kind, "", self.marker.marker_line,
                          self.marker.marker_line, self.marker.marker_line, meta)
            doc.results[bid] = BlockResult(bid, block.kind, content=content, state=state, reason=reason,
                                           text="" if heading else None)
            rendered.append(block)

        for stage, section in self.sections.items():
            add(stage, "heading_" + stage, heading=True)
            state, reason = self.status[stage]
            if stage in self.na:
                add(stage, stage + "_status", state=state, reason=reason, kind="text")
                continue
            if state not in ("PRESENT", "PASS"):
                add(stage, stage + "_status", state=state, reason=reason, kind="text")
                if stages[stage] or any(m.stage == stage and m.value is not None for m in self.materials.values()):
                    rendered[-1].meta["keep_next"] = "true"
            if not section.generate:
                rendered.extend(stages[stage])
                continue
            if stage == "overview":
                if self.identity:
                    item = self.identity
                    identity = " / ".join(str(v) for v in (item.id, item.name, item.revision) if v)
                    add(stage, "identity", kind="text")
                    doc.results["_packet_identity"].text = literal(identity)
                rows = [(self.sections[s].title, *self.status[s]) for s in self.sections if s != "overview"]
                add(stage, "overview", Table([Column("stage", "Section", align="left"),
                    Column("state", "State", align="left"), Column("reason", "Outstanding work", align="left")], rows))

            if stage == "compliance" and reqs:
                add(stage, "compliance", compliance_matrix(reqs))
            for name, material in self.materials.items():
                if material.stage != stage:
                    continue
                value = material.value
                if value is None:
                    if material.optional:
                        add(stage, name if name in SHEETS or name == "requirements" else "material_" + name,
                            state=material.state, reason=material.reason + " Optional supporting material.", kind="text")
                    continue
                if isinstance(value, Requirements):
                    value = requirements_table(value)
                elif isinstance(value, (Constants, Sheet)):
                    value = value.table()
                add(stage, name if name in SHEETS or name == "requirements" else "material_" + name,
                    value, kind="sources" if isinstance(value, Sources) else "table")
            if stage == "inputs":
                if reqs and reqs.all_variables():
                    add(stage, "controlled_variables", variables_table(reqs))
                if symbols is not None:
                    add(stage, "nomenclature", symbols)
            # The final matrix already includes these checks. Keep execution
            # and evidence, without printing a second verification panel.
            rendered.extend(b for b in stages[stage] if not (stage == "compliance" and b.kind == "verify" and reqs))
        doc.blocks = rendered

    @staticmethod
    def compliance_status(doc, reqs, work):
        if reqs is None:
            return "BLOCKED", "Requirements are not available for verification."
        else:
            _, failed, unverified = reqs.status()
            status = ("FAIL", f"{failed} failed verification(s).") if failed else (
                ("OPEN", f"{len(unverified)} requirement(s) need verification.") if unverified else ("PASS", "All local requirements verified."))
            if any(doc.results[b.id].state == "BLOCKED" or doc.results[b.id].failed for b in work):
                return "BLOCKED", "Verification work could not complete."
            return status

    @property
    def outstanding(self):
        return {stage: (state, reason) for stage, (state, reason) in self.status.items()
                if state in ("ERROR", "FAIL") or (self.sections[stage].required and state not in ("PRESENT", "PASS", "N/A"))}


def prepare(blocks, path):
    if any(b.kind == "packet" for b in blocks):
        return ComponentPacket(blocks, path)
    for block in blocks:
        if "stage" in block.meta:
            raise KipSyntaxError("stage= requires # %% packet component", path, block.marker_line)
    return None


def create_inputs(root):
    """Write the known blank sheets; never seed invented engineering data."""
    from openpyxl import Workbook
    from .sheets import save_workbook

    books = {}
    config = PacketConfig(root / "packet.toml")
    for name, spec in config.sheets.items():
        if name == "requirements":
            continue  # The scaffold uses the TOML alternative.
        if spec.path not in books:
            books[spec.path] = Workbook()
            books[spec.path].remove(books[spec.path].active)
        books[spec.path].create_sheet(spec.sheet).append(spec.columns)
    for path, book in books.items():
        save_workbook(book, root / path)
