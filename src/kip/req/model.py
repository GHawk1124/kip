"""Requirements, controlled variables, and the item tree.

A kip project is one *item* in a product breakdown -- a program, an assembly, a
sub-assembly or a component. Each item owns requirements, and requirements own
**controlled variables**: named, unit-carrying values that are the single source
of truth for a number, referenceable from any other design document.

That is what makes the flow-down real. A component's document does not restate
"the design load is 18 kN"; it reads ``reqs.P_design`` from the assembly that
levies the requirement, so changing the assembly changes every child document.

They live in ``input/requirements.xlsx``, one sheet per kind of row:

``Item``       one row: id, name, kind, revision, parent, parent_file, description
``Variables``  name, value, unit, description, source
``Inputs``     id, text, verification, parent, controls, rationale, note

``Item`` and ``Variables`` are optional; the item id then defaults to the
project folder's name. ``controls`` lists variable names separated by commas,
and ``parent_file`` names the parent item's workbook relative to this
project's folder (``../skid/input/requirements.xlsx``).
An older ``requirements.toml`` with ``[item]``, ``[vars.*]`` and ``[req.*]``
tables still loads; ``kip migrate`` converts it.
"""

from __future__ import annotations

import tomllib
from dataclasses import dataclass, field
from pathlib import Path

from ..units import fmt_quantity, ureg

__all__ = [
    "ControlledVar", "Requirement", "Item", "Check", "Requirements",
    "VERIFICATION_METHODS", "ITEM_KINDS", "RequirementsError",
]

def _resolve(path: "str | Path") -> Path:
    """Relative paths mean "next to this doc.py", whatever the working directory."""
    from ..doc.context import project_path
    return project_path(path).resolve()

#: Standard verification methods (as used in systems engineering practice).
VERIFICATION_METHODS = ("analysis", "test", "inspection", "demonstration",
                        "similarity")

#: Product-breakdown levels, coarse to fine.
ITEM_KINDS = ("program", "system", "assembly", "subassembly", "component")


class RequirementsError(Exception):
    """A malformed or inconsistent requirements file."""


@dataclass(frozen=True)
class ControlledVar:
    """A single controlled value. The authority for one number."""

    name: str
    value: float
    unit: str = ""
    description: str = ""
    source: str | None = None      # requirement id that levies it
    owner: str = ""                # item id that defines it

    @property
    def quantity(self):
        """The value as a pint quantity (dimensionless when no unit)."""
        return ureg.Quantity(self.value, self.unit or "dimensionless")

    def __str__(self) -> str:
        return fmt_quantity(self.quantity)


@dataclass
class Requirement:
    id: str
    text: str
    verification: str = "analysis"
    parent: str | None = None          # requirement this flows down from
    controls: list[str] = field(default_factory=list)
    rationale: str = ""
    owner: str = ""                    # item id
    note: str = ""

    def __post_init__(self) -> None:
        if self.verification not in VERIFICATION_METHODS:
            raise RequirementsError(
                f"{self.id}: verification={self.verification!r} is not one of "
                + ", ".join(VERIFICATION_METHODS)
            )


@dataclass
class Item:
    """One node of the product breakdown."""

    id: str
    name: str = ""
    kind: str = "component"
    parent: str | None = None
    parent_file: str | None = None
    revision: str = ""
    description: str = ""

    def __post_init__(self) -> None:
        if self.kind not in ITEM_KINDS:
            raise RequirementsError(
                f"item {self.id}: kind={self.kind!r} is not one of "
                + ", ".join(ITEM_KINDS)
            )


@dataclass
class Check:
    """One verification event binding a computed value to a requirement."""

    req_id: str
    value: object
    criterion: str
    passed: bool
    method: str = "analysis"
    evidence: str = ""              # block id or note
    detail: str = ""

    @property
    def status(self) -> str:
        return "PASS" if self.passed else "FAIL"

# criterion evaluation


_OPS = {
    ">=": lambda a, b: a >= b,
    "<=": lambda a, b: a <= b,
    "==": lambda a, b: a == b,
    "!=": lambda a, b: a != b,
    ">": lambda a, b: a > b,
    "<": lambda a, b: a < b,
}


def evaluate(value, criterion: str, variables: dict[str, ControlledVar] | None = None):
    """Evaluate ``value <op> bound``; returns ``(passed, rendered_bound)``.

    The bound may be a bare number, a quantity with units (``"<= 250 MPa"``) or
    the name of a controlled variable (``">= DF_min"``), which is what keeps a
    criterion tied to the requirement rather than to a literal typed twice.
    """
    text = criterion.strip()
    for op in (">=", "<=", "==", "!=", ">", "<"):
        if text.startswith(op):
            rhs = text[len(op):].strip()
            break
    else:
        raise RequirementsError(
            f"criterion {criterion!r} must start with one of "
            + ", ".join(_OPS)
        )

    variables = variables or {}
    if rhs in variables:
        bound = variables[rhs].quantity
    else:
        try:
            bound = ureg.Quantity(rhs)
        except Exception as e:
            raise RequirementsError(
                f"cannot parse the bound {rhs!r} in criterion {criterion!r}: {e}"
            ) from e

    # a bare number compared against a dimensionless quantity, and vice versa
    try:
        if isinstance(value, ureg.Quantity) and not isinstance(bound, ureg.Quantity):
            bound = ureg.Quantity(bound, value.units)
        elif isinstance(bound, ureg.Quantity) and not isinstance(value, ureg.Quantity):
            value = ureg.Quantity(value, bound.units)
        passed = bool(_OPS[op](value, bound))
    except Exception as e:
        raise RequirementsError(
            f"cannot compare {value!r} with {bound!r}: {e}"
        ) from e
    return passed, fmt_quantity(bound)

# workbooks

#: Where a project's requirements are looked for, in order.
REQUIREMENTS_FILES = ("input/requirements.xlsx", "requirements.toml")

ITEM_COLUMNS = ("id", "name", "kind", "revision", "parent", "parent_file", "description")
VARIABLE_COLUMNS = ("name", "value", "unit", "description", "source")
REQUIREMENT_COLUMNS = ("id", "text", "verification", "parent", "controls", "rationale", "note")


def _cell(value):
    """A blank cell is absent, not an empty string."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return value.strip() if isinstance(value, str) else value


def workbook_data(path: Path) -> dict:
    """Read a requirements workbook into the same shape as the TOML tables."""
    from openpyxl import load_workbook
    from ..authoring import read_records

    book = load_workbook(path, read_only=True)
    try:
        sheets = set(book.sheetnames)
    finally:
        book.close()
    if "Inputs" not in sheets:
        raise RequirementsError(f"{path}: expected an 'Inputs' sheet of requirements")

    def rows(sheet, required):
        if sheet not in sheets:
            return []
        records = read_records(path, sheet)
        for n, record in enumerate(records, start=2):
            for column in required:
                if _cell(record.get(column)) is None:
                    raise RequirementsError(f"{path} / {sheet} row {n}: {column} is blank")
        return [{k: _cell(v) for k, v in record.items()} for record in records]

    items = rows("Item", ("id",))
    if len(items) > 1:
        raise RequirementsError(f"{path} / Item: expected one row, found {len(items)}")
    folder = path.parent.parent if path.parent.name == "input" else path.parent
    item = {k: v for k, v in (items[0] if items else {"id": folder.name}).items() if v is not None}
    data: dict = {"item": item, "vars": {}, "req": {}}
    for row in rows("Variables", ("name", "value")):
        name = str(row["name"])
        if name in data["vars"]:
            raise RequirementsError(f"{path} / Variables: {name} appears twice")
        data["vars"][name] = {k: v for k, v in row.items() if k != "name" and v is not None}
    for row in rows("Inputs", ("id", "text")):
        rid = str(row["id"])
        if rid in data["req"]:
            raise RequirementsError(f"{path} / Inputs: {rid} appears twice")
        entry = {k: v for k, v in row.items() if k != "id" and v is not None}
        if "controls" in entry:
            entry["controls"] = [c.strip() for c in str(entry["controls"]).split(",") if c.strip()]
        data["req"][rid] = entry
    return data


# container


class Requirements:
    """Requirements and controlled variables for one item, plus its ancestors.

    Attribute access reads controlled variables as pint quantities, so a
    document says ``P = reqs.P_design`` rather than restating the number.
    Lookups walk up the parent chain, so a component transparently sees the
    values levied by the assembly above it.
    """

    def __init__(self, item: Item, requirements: dict[str, Requirement],
                 variables: dict[str, ControlledVar],
                 parent: "Requirements | None" = None,
                 path: Path | None = None):
        self.item = item
        self.requirements = requirements
        self.variables = variables
        self.parent = parent
        self.path = path
        self.checks: list[Check] = []

    # -- loading ---------------------------------------------------------
    @classmethod
    def load(cls, path: str | Path | None = None, _seen: set[Path] | None = None
             ) -> "Requirements":
        """Load ``input/requirements.xlsx`` (or an older ``requirements.toml``)."""
        if path is None:
            path = next((c for c in REQUIREMENTS_FILES if _resolve(c).exists()),
                        REQUIREMENTS_FILES[0])
        p = _resolve(path)
        _seen = _seen or set()
        if p in _seen:
            raise RequirementsError(f"circular parent_file chain at {p}")
        _seen.add(p)
        if not p.exists():
            raise RequirementsError(f"no requirements file at {p}")
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            return cls._from_dict(workbook_data(p), p, _seen)
        from ..doc.context import deprecated
        deprecated(f"{p.name} is an older form; requirements live in "
                   "input/requirements.xlsx (run kip migrate to convert it)")
        return cls._from_dict(tomllib.loads(p.read_text(encoding="utf-8")), p, _seen)

    @classmethod
    def loads(cls, text: str, path: Path | None = None) -> "Requirements":
        return cls._from_dict(tomllib.loads(text), path, set())

    @classmethod
    def _from_dict(cls, data: dict, path: Path | None, seen: set[Path]
                   ) -> "Requirements":
        raw_item = dict(data.get("item", {}) or {})
        if "id" not in raw_item:
            raise RequirementsError("requirements file needs [item] with an id")
        item = Item(
            id=str(raw_item["id"]), name=raw_item.get("name", ""),
            kind=raw_item.get("kind", "component"),
            parent=raw_item.get("parent"),
            parent_file=raw_item.get("parent_file"),
            revision=str(raw_item.get("revision", "")),
            description=raw_item.get("description", ""),
        )

        parent = None
        if item.parent_file and path is not None:
            # relative to the project folder, which for a workbook is above input/
            base = path.parent.parent if (path.suffix.lower() != ".toml"
                                          and path.parent.name == "input") else path.parent
            parent = cls.load((base / item.parent_file), _seen=seen)

        variables: dict[str, ControlledVar] = {}
        for name, raw in (data.get("vars", {}) or {}).items():
            if not isinstance(raw, dict) or "value" not in raw:
                raise RequirementsError(
                    f"[vars.{name}] needs at least a 'value'")
            variables[name] = ControlledVar(
                name=name, value=float(raw["value"]),
                unit=str(raw.get("unit", "")),
                description=raw.get("description", ""),
                source=raw.get("source"), owner=item.id,
            )

        requirements: dict[str, Requirement] = {}
        for rid, raw in (data.get("req", {}) or {}).items():
            raw = dict(raw)
            raw["text"] = str(raw.get("text", "")) if "text" in raw else None
            if raw["text"] is None:
                raise RequirementsError(f"[req.{rid}] needs 'text'")
            requirements[rid] = Requirement(
                id=rid, text=raw["text"],
                verification=raw.get("verification", "analysis"),
                parent=raw.get("parent"),
                controls=list(raw.get("controls", [])),
                rationale=raw.get("rationale", ""),
                note=raw.get("note", ""),
                owner=item.id,
            )

        obj = cls(item, requirements, variables, parent, path)
        obj._validate()
        return obj

    def _validate(self) -> None:
        for req in self.requirements.values():
            for name in req.controls:
                if self.var(name, strict=False) is None:
                    raise RequirementsError(
                        f"{req.id} controls {name!r}, which is not defined in "
                        f"[vars] of {self.item.id} or any parent"
                    )

    # -- variable access -------------------------------------------------
    def var(self, name: str, strict: bool = True) -> ControlledVar | None:
        if name in self.variables:
            return self.variables[name]
        if self.parent is not None:
            found = self.parent.var(name, strict=False)
            if found is not None:
                return found
        if strict:
            raise AttributeError(
                f"no controlled variable {name!r} in {self.item.id}"
                + (f" or its ancestors" if self.parent else "")
            )
        return None

    def all_variables(self) -> dict[str, ControlledVar]:
        """Every visible variable, nearest definition winning."""
        out: dict[str, ControlledVar] = {}
        if self.parent is not None:
            out.update(self.parent.all_variables())
        out.update(self.variables)
        return out

    def all_requirements(self) -> dict[str, Requirement]:
        out: dict[str, Requirement] = {}
        if self.parent is not None:
            out.update(self.parent.all_requirements())
        out.update(self.requirements)
        return out

    def __getattr__(self, name: str):
        if name.startswith("_"):
            raise AttributeError(name)
        return self[name]

    def __getitem__(self, name: str):
        from ..quantities import controlled_entry, record_read
        variable = self.var(name)
        record_read(self, controlled_entry(variable))
        return variable.quantity

    def __contains__(self, name: str) -> bool:
        return self.var(name, strict=False) is not None

    def lineage(self) -> list[Item]:
        """This item and its ancestors, root first."""
        chain = [self.item]
        node = self.parent
        while node is not None:
            chain.append(node.item)
            node = node.parent
        return list(reversed(chain))

    # -- verification ----------------------------------------------------
    def verify(self, req_id: str, value, criterion: str,
               evidence: str = "", note: str = "") -> Check:
        """Bind a computed value to a requirement and record pass/fail."""
        reqs = self.all_requirements()
        if req_id not in reqs:
            raise RequirementsError(
                f"cannot verify unknown requirement {req_id!r}; defined: "
                + ", ".join(sorted(reqs)) or "(none)"
            )
        req = reqs[req_id]
        passed, bound = evaluate(value, criterion, self.all_variables())
        check = Check(
            req_id=req_id, value=value, criterion=criterion, passed=passed,
            method=req.verification, evidence=evidence, detail=note,
        )
        self.checks.append(check)
        return check

    def status(self, inherited: bool = False) -> tuple[int, int, list[str]]:
        """``(passed, failed, unverified_requirement_ids)``.

        Only this item's own requirements count as unverified by default: a
        parent's requirement is discharged by the parent's own document, not by
        every child that inherits its controlled variables.
        """
        scope = self.all_requirements() if inherited else self.requirements
        verified = {c.req_id for c in self.checks}
        passed = sum(1 for c in self.checks if c.passed)
        failed = sum(1 for c in self.checks if not c.passed)
        unverified = sorted(set(scope) - verified)
        return passed, failed, unverified

    @property
    def compliant(self) -> bool:
        passed, failed, unverified = self.status()
        return failed == 0 and not unverified
