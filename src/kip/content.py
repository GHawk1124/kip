"""Rich content a block can produce: figures, tables, drawings and sources.

A block binds one of these objects and the renderer picks it up, so authoring
stays ordinary Python::

    # %% kip.plot id=curve
    fig = Figure(xlabel="Deflection (mm)", ylabel="Load (kN)")
    fig.line(defl, load, label="Test", mark="o")

Everything here is *data*; nothing renders until the emitter walks it.  That
keeps the objects importable without Typst and testable without compiling.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from .units import fmt_quantity, ureg

__all__ = [
    "Figure", "Series", "Table", "Column", "Drawing", "Source", "Sources",
    "strip_units", "RichContent", "Symbol", "Math", "nomenclature", "plot",
]


def strip_units(values: Iterable[Any], unit: str | None = None) -> list[float]:
    """Convert a sequence of quantities to plain floats for plotting.

    If ``unit`` is given every element is converted to it first, so a series
    mixing mm and m plots consistently rather than silently mis-scaling.
    """
    out: list[float] = []
    for v in values:
        if isinstance(v, ureg.Quantity):
            out.append(float(v.to(unit).magnitude if unit else v.magnitude))
        else:
            out.append(float(v))
    return out


def _infer_unit(values: Sequence[Any]) -> str | None:
    for v in values:
        if isinstance(v, ureg.Quantity):
            return f"{v.units:~P}"
    return None

# figures


@dataclass
class Series:
    """One plotted series."""

    x: list[float]
    y: list[float]
    label: str | None = None
    kind: str = "line"           # line | scatter | bar
    mark: str | None = None      # o, square, triangle, ...
    dash: str | None = None      # dashed, dotted, ...
    color: str | None = None


@dataclass
class Figure:
    """A 2-D plot rendered as vector graphics by lilaq.

    Axis labels pick up units automatically when the data are pint quantities,
    so ``fig.line(defl_mm, load_kn)`` is labelled ``Load (kN)`` without the
    author restating it.
    """

    xlabel: str | None = None
    ylabel: str | None = None
    title: str | None = None
    width: float = 120.0          # mm
    height: float = 70.0          # mm
    xscale: str = "linear"        # linear | log
    yscale: str = "linear"
    grid: bool = True
    legend: bool = True
    series: list[Series] = field(default_factory=list)
    xlim: tuple[float, float] | None = None
    ylim: tuple[float, float] | None = None

    def _add(self, kind, x, y, label, mark, dash, color, xunit, yunit) -> "Figure":
        x, y = list(x), list(y)
        if len(x) != len(y):
            raise ValueError("plot x and y must have the same length")
        xu = xunit or _infer_unit(x)
        yu = yunit or _infer_unit(y)
        if self.xlabel and xu and "(" not in self.xlabel:
            self.xlabel = f"{self.xlabel} ({xu})"
        if self.ylabel and yu and "(" not in self.ylabel:
            self.ylabel = f"{self.ylabel} ({yu})"
        self.series.append(Series(
            x=strip_units(x, xu), y=strip_units(y, yu),
            label=label, kind=kind, mark=mark, dash=dash, color=color,
        ))
        return self

    def line(self, x, y, label=None, mark=None, dash=None, color=None,
             xunit=None, yunit=None) -> "Figure":
        return self._add("line", x, y, label, mark, dash, color, xunit, yunit)

    def scatter(self, x, y, label=None, mark="o", color=None,
                xunit=None, yunit=None) -> "Figure":
        return self._add("scatter", x, y, label, mark, None, color, xunit, yunit)

    def bar(self, x, y, label=None, color=None, xunit=None, yunit=None) -> "Figure":
        return self._add("bar", x, y, label, None, None, color, xunit, yunit)

    def of(self, expr, symbol, start, stop, n: int = 100, label=None, **kw) -> "Figure":
        """Plot a SymPy expression over a range -- verification by analysis."""
        import sympy as sp

        f = sp.lambdify(symbol, expr, "math")
        step = (stop - start) / max(1, n - 1)
        xs = [start + i * step for i in range(n)]
        ys = []
        for x in xs:
            try:
                ys.append(float(f(x)))
            except (ValueError, ZeroDivisionError, OverflowError):
                ys.append(float("nan"))
        return self.line(xs, ys, label=label or str(expr), **kw)

# tables


@dataclass(frozen=True)
class Symbol:
    """An engineering identifier with Greek letters and subscripts."""

    name: str

    def __str__(self) -> str:
        return self.name


@dataclass(frozen=True)
class Math:
    """Explicit inline Typst math, e.g. Math('sigma_y / 2')."""

    body: str

    def __str__(self) -> str:
        return self.body


@dataclass
class Column:
    key: str
    title: str | Symbol | Math | None = None
    unit: str | None = None       # convert quantities to this before display
    align: str = "right"
    precision: int = 3
    format: str | None = None     # e.g. "{:.1%}"
    math: bool = False            # identifier strings render as math symbols


@dataclass
class Table:
    """Tabular data, rendered as a Typst table and optionally exported to xlsx.

    ``max_rows`` truncates the *rendered* table while the exported workbook
    keeps every row -- a 400-row sweep belongs in the spreadsheet, not in the
    middle of a design note.
    """

    columns: list[Column | str]
    rows: list[Sequence[Any]] = field(default_factory=list)
    title: str | None = None
    caption: str | None = None
    xlsx: str | None = None       # filename, written next to the PDF
    max_rows: int | None = None
    highlight: dict[int, str] = field(default_factory=dict)  # row -> "ok"|"fail"
    zebra: bool = True
    total_row: Sequence[Any] | None = None

    def __post_init__(self) -> None:
        self.columns = [
            c if isinstance(c, Column) else Column(key=str(c), title=str(c))
            for c in self.columns
        ]

        self.rows = list(self.rows)
        if not self.columns:
            raise ValueError("a table needs at least one column")
        for i, row in enumerate(self.rows):
            if len(row) != len(self.columns):
                raise ValueError(f"table row {i + 1}: expected {len(self.columns)} cells, got {len(row)}")
        if self.total_row is not None and len(self.total_row) != len(self.columns):
            raise ValueError("table total row must match the column count")
        if self.max_rows is not None and self.max_rows < 0:
            raise ValueError("max_rows must be non-negative")

    @classmethod
    def from_records(cls, records, **options) -> "Table":
        """Dictionary keys become headers, preserving insertion order."""
        records = list(records)
        if not records:
            raise ValueError("from_records needs at least one record; use Table for an empty table")
        keys = list(records[0])
        if any(set(record) != set(keys) for record in records):
            raise ValueError("all table records must have the same keys")
        return cls(columns=keys, rows=[[record[k] for k in keys] for record in records], **options)

    @property
    def headers(self) -> list:
        out = []
        for c in self.columns:
            title = c.title or c.key
            out.append(f"{title} ({c.unit})" if c.unit and "(" not in str(title) else title)
        return out

    def cell_text(self, value: Any, col: Column) -> str:
        if value is None:
            return ""
        if isinstance(value, ureg.Quantity):
            if col.unit:
                value = value.to(col.unit)
                return f"{value.magnitude:.{col.precision}f}"
            return fmt_quantity(value, col.precision)
        if col.format:
            try:
                return col.format.format(value)
            except (ValueError, KeyError, IndexError):
                pass
        if isinstance(value, float):
            return f"{value:.{col.precision}f}"
        return str(value)

    def display_rows(self) -> tuple[list[list[str]], int]:
        """Formatted rows for the PDF, plus the count that was hidden."""
        shown = self.rows if self.max_rows is None else self.rows[: self.max_rows]
        hidden = len(self.rows) - len(shown)
        out = [
            [self.cell_text(v, c) for v, c in zip(row, self.columns)]
            for row in shown
        ]
        return out, hidden

    def raw_value(self, value: Any, col: Column) -> Any:
        """Native value for the spreadsheet: numbers stay numbers."""
        if isinstance(value, ureg.Quantity):
            return float(value.to(col.unit).magnitude if col.unit
                         else value.magnitude)
        from sympy import Basic
        return str(value) if isinstance(value, (Symbol, Math, Basic)) else value

    def to_xlsx(self, path: str | Path) -> Path:
        """Write the full table (never truncated) as a real workbook."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = (self.title or "Data")[:31]

        head_fill = PatternFill("solid", fgColor="EDE6D3")
        head_font = Font(bold=True, color="1A1714")
        for j, header in enumerate(self.headers, start=1):
            cell = ws.cell(row=1, column=j, value=str(header))
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for i, row in enumerate(self.rows, start=2):
            for j, (value, col) in enumerate(zip(row, self.columns), start=1):
                cell = ws.cell(row=i, column=j, value=self.raw_value(value, col))
                cell.alignment = Alignment(horizontal=col.align, vertical="center")

        if self.total_row:
            r = len(self.rows) + 2
            for j, (value, col) in enumerate(zip(self.total_row, self.columns), start=1):
                cell = ws.cell(row=r, column=j, value=self.raw_value(value, col))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal=col.align, vertical="center")

        for j, header in enumerate(self.headers, start=1):
            width = max(len(str(header)) + 2,
                        *(len(self.cell_text(row[j - 1], self.columns[j - 1])) + 2
                          for row in self.rows)) if self.rows else len(str(header)) + 2
            ws.column_dimensions[get_column_letter(j)].width = min(32, width)
        ws.freeze_panes = "A2"

        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)
        return p

# drawings


@dataclass
class Drawing:
    """A vector drawing: CeTZ ``body`` or SVG bytes supplied by ``kip.cad``.

    ``width`` and ``height`` bound the displayed illustration. Attachments such
    as planar DXFs are exported beside the PDF and linked beneath the view.
    """

    body: str = ""
    width: float | None = None    # mm; None lets CeTZ size itself
    height: float | None = None
    caption: str | None = None
    length: str = "1cm"           # CeTZ unit length
    svg: bytes | None = None       # vector CAD projection, instead of CeTZ
    attachments: dict[str, bytes] = field(default_factory=dict)

# sources


@dataclass
class Source:
    """A citable reference. Rendered as a clickable, formatted link."""

    title: str
    url: str | None = None
    author: str | None = None
    section: str | None = None
    year: int | str | None = None
    publisher: str | None = None
    note: str | None = None

    def short(self) -> str:
        """Compact in-line form, e.g. 'ASME B31.3 §302.3.5'."""
        bits = [self.title]
        if self.section:
            bits.append(self.section)
        return " ".join(bits)

    def full(self) -> str:
        """Reference-list form."""
        bits: list[str] = []
        if self.author:
            bits.append(self.author)
        bits.append(self.title)
        if self.publisher:
            bits.append(self.publisher)
        if self.year:
            bits.append(str(self.year))
        if self.section:
            bits.append(self.section)
        out = ", ".join(b for b in bits if b)
        if self.note:
            out += f" — {self.note}"
        return out


class Sources(dict):
    """Mapping of citation key -> :class:`Source`, cited with ``@src:key``."""

    def __init__(self, mapping: dict[str, Source | dict] | None = None, **kw):
        super().__init__()
        for k, v in {**(mapping or {}), **kw}.items():
            self[k] = v if isinstance(v, Source) else Source(**v)


def nomenclature(entries: dict[str, str | tuple[str, str]], **options) -> Table:
    """Symbol -> description or (description, unit), rendered as a math table."""
    rows = []
    for name, entry in entries.items():
        description, unit = (entry, "-") if isinstance(entry, str) else entry
        rows.append((Symbol(name), description, unit))
    return Table([Column("symbol", "Symbol", align="left"),
                  Column("definition", "Definition", align="left"),
                  Column("unit", "Unit", align="left")], rows, **options)


def plot(x, y, *, label=None, mark=None, dash=None, color=None,
         xunit=None, yunit=None, **options) -> Figure:
    """Create a line plot in one call; chain .line(), .scatter() or .bar()."""
    return Figure(**options).line(x, y, label=label, mark=mark, dash=dash,
                                 color=color, xunit=xunit, yunit=yunit)


#: Types a block may bind for the renderer to pick up.
RichContent = (Figure, Table, Drawing, Sources)
