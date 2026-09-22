from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook, Workbook
from typer.testing import CliRunner

from kip.cli import app
from kip.doc import build
from kip.doc.kernel import execute, ExecutionError
from kip.doc.loader import KipSyntaxError
from kip.doc.validate import ValidationError
from kip.packet import SHEETS, STAGES, create_inputs
from kip.render import compile_pdf, emit
from kip.scaffold import create_project


def write_rows(root, name, rows):
    spec = SHEETS[name]
    path = root / spec.path
    book = load_workbook(path)
    try:
        for row in rows:
            book[spec.sheet].append([row.get(k) for k in spec.columns])
        book.save(path)
    finally:
        book.close()


def populated(root):
    create_inputs(root)
    write_rows(root, "constants", [{"key": "P", "value": 12, "unit": "kN", "description": "Load"}])
    (root / "requirements.toml").write_text('''[item]
id = "PART-001"
name = "Bracket"
[req.REQ-001]
text = "Load shall not exceed 20 kN."
verification = "analysis"
''')


def test_blank_component_builds_with_all_slots_open_and_check_fails(tmp_path):
    create_project(tmp_path / "part", template="component", source=str(Path.cwd()))
    path = tmp_path / "part/doc.py"
    doc = build(path)
    assert not doc.errors
    assert list(doc.packet.status) == list(STAGES)
    assert doc.packet.status["inputs"][0] == "OPEN"
    assert doc.packet.status["compliance"][0] == "BLOCKED"
    pdf = pymupdf.open(stream=compile_pdf(emit(doc)), filetype="pdf")
    text = "".join(p.get_text() for p in pdf)
    assert "PART" in text and "OPEN" in text and "BLOCKED" in text
    assert [title for level, title, page in pdf.get_toc() if level == 1] == [f"{i} {doc.packet.sections[s].title}" for i, s in enumerate(STAGES, 1)]
    assert "4 Preliminary Analysis" in text and "5 Sizing" in text
    runner = CliRunner()
    assert runner.invoke(app, ["check", str(path), "--render"]).exit_code == 1
    assert runner.invoke(app, ["build", str(path)]).exit_code == 0
    shown = runner.invoke(app, ["show", str(path)])
    assert shown.exit_code == 0, shown.output
    assert "OPEN" in shown.output


def test_missing_inputs_block_only_dependent_work_and_refresh_on_rebuild(tmp_path):
    (tmp_path / "analysis.py").write_text('''from kip import *
@calculation
def double(c):
    P = c.P
    # equations
    twice = 2 * P
''')
    source = '''from kip import *
import analysis
# %% packet component
# %% calculation "Initial sizing" stage=sizing
result = analysis.double(C)
# %% table "Inputs"
inputs_table(result)
# %% calc "Independent check"
known = 2 + 2
'''
    path = tmp_path / "doc.py"
    doc = build(path=path, source=source)
    assert doc.results["initial_sizing"].state == "BLOCKED"
    assert doc.results["inputs"].state == "BLOCKED"
    assert doc.results["independent_check"].values["known"] == 4
    populated(tmp_path)
    execute(doc)
    assert doc.results["initial_sizing"].state == "PRESENT"
    assert doc.namespace["result"].twice.magnitude == 24
    assert doc.packet.status["inputs"][0] == "PRESENT"
    book = load_workbook(tmp_path / "input/constants.xlsx")
    book.active["B2"] = 15
    book.save(tmp_path / "input/constants.xlsx")
    book.close()
    execute(doc)
    assert doc.namespace["result"].twice.magnitude == 30


def test_packet_order_is_independent_of_authored_order_and_na_is_explicit(tmp_path):
    populated(tmp_path)
    source = '''from kip import *
# %% packet component na="manufacturing,test"
# %% calc "Analysis work"
b = 3 + 4
# %% text "Design rationale" stage=design
"""Design decision."""
# %% sizing "Sizing work"
a = 1 + 2
'''
    doc = build(path=tmp_path / "doc.py", source=source)
    ids = [b.id for b in doc.ordered_blocks()]
    assert ids.index("sizing_work") < ids.index("design_rationale") < ids.index("analysis_work")
    assert doc.packet.status["manufacturing"][0] == "N/A"
    assert "manufacturing" not in doc.packet.outstanding
    assert doc.packet.status["compliance"][0] == "OPEN"


@pytest.mark.parametrize("failure", ["unit", "headers", "duplicate"])
def test_invalid_data_is_an_error_not_an_open_slot(tmp_path, failure):
    populated(tmp_path)
    path = tmp_path / "input/constants.xlsx"
    book = load_workbook(path)
    if failure == "unit":
        book.active["C2"] = "bananas_not_units"
    elif failure == "headers":
        book.active["A1"] = "wrong_key"
    else:
        book.active.append(["P", 15, "kN"])
    book.save(path)
    book.close()
    with pytest.raises(ValidationError):
        build(path=tmp_path / "doc.py", source="# %% packet component")


def test_completed_calculation_does_not_automatically_pass_compliance(tmp_path):
    populated(tmp_path)
    source = '''from kip import *
# %% packet component na="preliminary,sizing,design,manufacturing,test"
# %% calc "Load"
P = 12*kN
'''
    doc = build(path=tmp_path / "doc.py", source=source)
    assert doc.packet.status["compliance"][0] == "OPEN"
    checked = build(path=tmp_path / "doc.py", source=source + '''
# %% verify "Verification"
reqs.verify("REQ-001", P, "<= 20 kN", evidence="load")
''')
    assert checked.packet.status["compliance"][0] == "PASS"
    assert checked.results["_packet_compliance"].content.rows[0][-1] == "PASS"
    assert not checked.packet.outstanding


def test_test_plan_without_results_stays_open(tmp_path):
    populated(tmp_path)
    write_rows(tmp_path, "tests", [{"id": "T-1", "requirement": "REQ-001", "procedure": "Apply load", "criterion": "<= 20 kN"}])
    doc = build(path=tmp_path / "doc.py", source="# %% packet component")
    assert doc.packet.status["test"][0] == "OPEN"
    assert len(doc.results["_packet_tests"].content.rows) == 1
    assert doc.packet.status["compliance"][0] == "OPEN"


@pytest.mark.parametrize("source", [
    '# %% packet component na="test"\n# %% text "Test" stage=test\n"Work"',
    '# %% packet component\n# %% text "Typo" stage=analysiss\n"Work"',
    '# %% packet component\n# %% packet another',
    'C = 1\n# %% packet component',
])
def test_ambiguous_packet_authoring_is_rejected(source):
    with pytest.raises(KipSyntaxError):
        build(source=source)


def test_existing_documents_keep_strict_missing_file_behavior(tmp_path):
    with pytest.raises(ExecutionError):
        build(path=tmp_path / "doc.py", source='from kip import *\nC = Constants.load()\n# %% text "Scope"\n"Prose"')


def test_missing_drawing_is_open_but_invalid_code_still_fails(tmp_path):
    source = '''from kip import *
# %% packet component
# %% draw "Front view"
Drawing.load("assets/front.svg")
'''
    doc = build(path=tmp_path / "doc.py", source=source)
    assert doc.results["front_view"].state == "OPEN"
    assert doc.packet.status["design"][0] == "OPEN"
    compile_pdf(emit(doc))
    with pytest.raises(ExecutionError):
        build(path=tmp_path / "doc.py", source=source.replace('Drawing.load("assets/front.svg")', "unknown_function()"))


def test_requirement_workbook_is_an_alternative_not_a_second_authority(tmp_path):
    create_inputs(tmp_path)
    book = Workbook()
    book.active.title = "Inputs"
    book.active.append(["id", "text", "verification"])
    book.active.append(["REQ-1", "Meet the load limit", "analysis"])
    book.save(tmp_path / "input/requirements.xlsx")
    book.close()
    doc = build(path=tmp_path / "doc.py", source="# %% packet component")
    assert doc.namespace["reqs"].requirements["REQ-1"].text == "Meet the load limit"
    (tmp_path / "requirements.toml").write_text('[item]\nid="PART"')
    with pytest.raises(ValidationError, match="choose requirements.toml"):
        build(path=tmp_path / "doc.py", source="# %% packet component")


def test_test_requirement_links_and_na_contradictions_are_validated(tmp_path):
    populated(tmp_path)
    write_rows(tmp_path, "tests", [{"id": "T-1", "requirement": "TYPO", "procedure": "Apply load", "criterion": "<= 20 kN"}])
    with pytest.raises(ValidationError, match="unknown requirement"):
        build(path=tmp_path / "doc.py", source="# %% packet component")
    with pytest.raises(ValidationError, match="N/A but tests contains data"):
        build(path=tmp_path / "doc.py", source='# %% packet component na="test"')


def test_failed_verification_is_visible_in_the_final_overview_and_matrix(tmp_path):
    populated(tmp_path)
    source = '''from kip import *
# %% packet component
# %% verify "Load check"
reqs.verify("REQ-001", 25*kN, "<= 20 kN", evidence="Bench result")
'''
    doc = build(path=tmp_path / "doc.py", source=source)
    assert doc.packet.status["compliance"][0] == "FAIL"
    assert doc.results["_packet_compliance"].content.rows[0][-1] == "FAIL"
    assert doc.results["_packet_overview"].content.rows[-1][1] == "FAIL"
    pdf = pymupdf.open(stream=compile_pdf(emit(doc)), filetype="pdf")
    assert "FAIL" in "".join(p.get_text() for p in pdf)


def test_packet_nomenclature_conflicts_are_actionable_validation_errors(tmp_path):
    populated(tmp_path)
    source = '''from kip import *
# %% packet component
# %% inputs "Local definition"
P = 1 * mm
'''
    with pytest.raises(ValidationError, match="incompatible units"):
        build(path=tmp_path / "doc.py", source=source)


def test_missing_expected_sheet_is_open_but_present_bad_headers_are_errors(tmp_path):
    populated(tmp_path)
    path = tmp_path / "input/references.xlsx"
    book = load_workbook(path)
    book.remove(book["Documents"])
    book.save(path)
    book.close()
    doc = build(path=tmp_path / "doc.py", source="# %% packet component")
    assert doc.packet.materials["documents"].state == "OPEN"
    assert "Missing sheet" in doc.packet.materials["documents"].reason
    assert not doc.errors


def test_workbook_definition_takes_precedence_over_automatic_cell_title(tmp_path):
    populated(tmp_path)
    doc = build(path=tmp_path / "doc.py", source='''from kip import *
# %% packet component
# %% sizing "Initial sizing"
P = 12 * kN
''')
    rows = {row[0].name: row[1] for row in doc.results["_packet_nomenclature"].content.rows}
    assert rows["P"] == "Load"


def test_custom_order_titles_defaults_and_optional_sections(tmp_path):
    (tmp_path / "packet.toml").write_text('''
order = ["review", "preliminary", "sizing", "appendix"]
[defaults]
text = "review"
[sections.review]
title = "Design Review"
[sections.appendix]
required = false
''')
    doc = build(path=tmp_path / "doc.py", source='''
# %% packet component
# %% sizing "Size"
size = 4 + 5
# %% text "Decision"
"Proceed to detailed design."
# %% preliminary "Feasibility"
estimate = 2 + 3
''')
    assert list(doc.packet.status) == ["review", "preliminary", "sizing", "appendix"]
    ids = [b.id for b in doc.ordered_blocks()]
    assert ids.index("decision") < ids.index("feasibility") < ids.index("size")
    assert doc.packet.status["appendix"][0] == "OPEN"
    assert not doc.packet.outstanding
    pdf = pymupdf.open(stream=compile_pdf(emit(doc)))
    assert [title for _, title, _ in pdf.get_toc()] == [
        "1 Design Review", "2 Preliminary Analysis", "3 Sizing", "4 Appendix"]


def test_custom_sheet_paths_tabs_columns_and_required_values(tmp_path):
    (tmp_path / "packet.toml").write_text('''
order = ["inputs", "risks"]
[sheets.constants]
path = "data/loads.xlsx"
sheet = "Loads"
columns = ["key", "value", "unit"]
[sheets.hazards]
path = "data/risks.xlsx"
stage = "risks"
columns = ["hazard", "mitigation"]
required = ["hazard", "mitigation"]
''')
    create_inputs(tmp_path)
    for file, row in (("loads", ["P", 12, "kN"]), ("risks", ["Overload", "Limit force"])):
        path = tmp_path / f"data/{file}.xlsx"
        book = load_workbook(path)
        book.active.append(row)
        book.save(path)
        book.close()
    source = '# %% packet component\n# %% table "Load" stage=inputs\nC.table()'
    doc = build(path=tmp_path / "doc.py", source=source)
    assert doc.namespace["C"].P.magnitude == 12
    assert doc.results["_packet_material_hazards"].content.rows == [["Overload", "Limit force"]]
    assert not doc.packet.outstanding
    book = load_workbook(tmp_path / "data/risks.xlsx")
    book.active["B2"] = None
    book.save(tmp_path / "data/risks.xlsx")
    book.close()
    with pytest.raises(ValidationError, match="mitigation is empty"):
        build(path=tmp_path / "doc.py", source=source)


def test_generated_tables_can_be_replaced_without_disabling_input_validation(tmp_path):
    populated(tmp_path)
    (tmp_path / "packet.toml").write_text('''
order = ["inputs"]
[sections.inputs]
generate = false
''')
    source = '# %% packet component\n# %% table "Selected input" stage=inputs\nC.table()'
    doc = build(path=tmp_path / "doc.py", source=source)
    assert "_packet_constants" not in doc.results and "_packet_nomenclature" not in doc.results
    assert doc.namespace["C"].P.magnitude == 12
    assert not doc.packet.outstanding
    book = load_workbook(tmp_path / "input/constants.xlsx")
    book.active["C2"] = "not_a_unit"
    book.save(tmp_path / "input/constants.xlsx")
    book.close()
    with pytest.raises(ValidationError):
        build(path=tmp_path / "doc.py", source=source)


def test_disabled_readers_allow_authored_replacements_and_custom_config_path(tmp_path):
    (tmp_path / "custom.toml").write_text('''
order = ["review", "inputs"]
na = ["review"]
[sheets.constants]
enabled = false
''')
    doc = build(path=tmp_path / "doc.py", source='''
from types import SimpleNamespace
C = SimpleNamespace(load=12)
# %% packet component config="custom.toml"
# %% calc "Input" stage=inputs
P = C.load
''')
    assert doc.packet.status["review"][0] == "N/A"
    assert doc.namespace["P"] == 12
    assert not doc.packet.outstanding


@pytest.mark.parametrize("config", [
    'order = ["inputs", "inputs"]',
    'order = []',
    'order = ["inputs"]\n[sections.typo]\ncolumns=2',
    '[sections.analysis]\ncolumns=3',
    '[sections.analysis]\ncolumns=true',
    '[sections.analysis]\ngenerate="false"',
    '[sections.analysis]\ntitel="Oops"',
    '[defaults]\ncalc="typo"',
    '[sheets.constants]\nstage="typo"',
    '[sheets.constants]\ncolumns=["value"]',
    '[sheets.extra]\npath="data.xlsx"',
    'na = ["typo"]',
])
def test_invalid_packet_configuration_points_to_its_file(tmp_path, config):
    (tmp_path / "packet.toml").write_text(config)
    with pytest.raises(KipSyntaxError, match="packet.toml"):
        build(path=tmp_path / "doc.py", source="# %% packet component")


def test_removed_sections_reject_authored_work_instead_of_dropping_it(tmp_path):
    (tmp_path / "packet.toml").write_text('order = ["preliminary"]')
    with pytest.raises(KipSyntaxError, match="removed stage 'analysis'"):
        build(path=tmp_path / "doc.py", source='# %% packet component\n# %% calc "Work"\na=1+2')


@pytest.mark.parametrize("columns", [1, 2])
def test_packet_columns_can_switch_mid_section_and_reset_at_next_section(tmp_path, columns):
    from kip.render.layout import Layout, PageSpec
    (tmp_path / "packet.toml").write_text('''
order = ["preliminary", "design"]
[sections.preliminary]
columns = 2
''')
    doc = build(path=tmp_path / "doc.py", source='''
# %% packet component
# %% text "Narrow" stage=preliminary
"Two column working."
# %% text "Wide" stage=preliminary columns=1
"Full width detail."
# %% text "Restored" stage=preliminary columns=default
"Document default."
# %% text "Local" stage=preliminary columns=2
"Local passage."
# %% text "Design" stage=design
"Next section uses its default."
''')
    files = emit(doc, Layout(page=PageSpec(columns=columns)))
    text = files["main.typ"].decode()
    # The second stage heading restores the document default even after a local override.
    assert next(b for b in doc.blocks if b.id == "_packet_heading_design").meta["columns"] == "default"
    start = text.index('id: "wide"')
    restored = text.index('id: "restored"')
    local = text.index('id: "local"')
    assert ("#columns(2" in text[start:restored]) == (columns == 2)
    assert ("#columns(2" in text[restored:local]) == (columns == 1)
    pdf = pymupdf.open(stream=compile_pdf(files))
    output = "".join(p.get_text() for p in pdf)
    assert "Next section uses its default." in output
    for page in pdf:
        assert all(0 <= x0 < x1 <= page.rect.width for x0, _, x1, *_ in page.get_text("words"))


def test_new_cli_can_set_two_column_default(tmp_path):
    from kip.render.layout import Layout
    root = tmp_path / "two-column"
    result = CliRunner().invoke(app, ["new", str(root), "--template", "component", "--columns", "2", "--no-sync"])
    assert result.exit_code == 0, result.output
    assert "run_document(__file__, title=\"Two Column\", columns=2)" in (root / "doc.py").read_text()
    assert not (root / "layout.toml").exists()
    doc = build(root / "doc.py")
    pdf = pymupdf.open(stream=compile_pdf(emit(doc)))
    assert "4 Preliminary Analysis" in "".join(p.get_text() for p in pdf)


def test_packet_restores_identity_header_and_renders_checks_once(tmp_path):
    populated(tmp_path)
    (tmp_path / "packet.toml").write_text('order = ["overview", "requirements", "compliance"]')
    doc = build(path=tmp_path / "doc.py", source='''from kip import *
# %% packet component
# %% verify "Load check" stage=compliance
reqs.verify("REQ-001", 12*kN, "<= 20*kN", evidence="Measured load")
''')
    files = emit(doc)
    main = files["main.typ"].decode()
    assert "PART-001" in main and "Document" in main
    assert "#kip-verify(" not in main
    assert "id: \"load_check\"" not in main
    pdf = pymupdf.open(stream=compile_pdf(files))
    text = "".join(p.get_text() for p in pdf)
    assert text.count("Measured load") == 1
    assert "Assessment" in text and "Status" in text
    assert all("PART-001" in page.get_text() for page in pdf)
