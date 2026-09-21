"""Spreadsheet inputs: constants carry units, sheets render themselves."""
import pytest
from openpyxl import Workbook

from kip import Constants, Sheet, Sources, calculation, mm, ureg
from kip.doc import build
from kip.doc.kernel import ExecutionError
from kip.units import fmt_quantity


def workbook(tmp_path, name, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    path = tmp_path / name
    wb.save(path)
    return path


@pytest.fixture
def constants_file(tmp_path):
    return workbook(
        tmp_path, "constants.xlsx",
        ["key", "value", "unit", "description", "basis", "symbol"],
        [("rho_w", 997.99, "kg/m3", "Water density", "IAPWS", ""),
         ("d_face", 58.0, "mm", "Screen diameter", "Design", ""),
         ("Cv_f", 6e7, "1/m", "Viscous resistance", "Assumed", "C_vf"),
         ("growth", 1.2, "1", "Reserve factor", "Assumption", "")],
    )


# constants


def test_constants_carry_the_unit_the_sheet_declares(constants_file):
    C = Constants.load(constants_file)
    assert C.rho_w == ureg.Quantity(997.99, "kg/m**3")
    assert C["d_face"].to(mm).magnitude == pytest.approx(58.0)
    assert C.growth.dimensionless
    assert C.value("d_face", mm) == pytest.approx(58.0)


def test_spreadsheet_exponent_shorthand_parses(constants_file):
    # "kg/m3" is how a unit is written in a workbook; pint wants kg/m**3.
    assert Constants.load(constants_file).rho_w.units == ureg.kg / ureg.m ** 3


def test_a_bad_unit_fails_at_the_row_that_declares_it(tmp_path):
    path = workbook(tmp_path, "c.xlsx", ["key", "value", "unit"],
                    [("x", 1.0, "furlongs_per_fortnight")])
    with pytest.raises(ValueError, match="row 2: x:"):
        Constants.load(path)


def test_non_numeric_and_duplicate_keys_are_rejected(tmp_path):
    bad = workbook(tmp_path, "a.xlsx", ["key", "value"], [("x", "lots")])
    with pytest.raises(ValueError, match="numeric value"):
        Constants.load(bad)
    dup = workbook(tmp_path, "b.xlsx", ["key", "value"], [("x", 1), ("x", 2)])
    with pytest.raises(ValueError, match="duplicate key"):
        Constants.load(dup)


def test_constants_can_be_generated_programmatically(constants_file):
    C = Constants.load(constants_file)
    C.add("rho_hot", 971.8, "kg/m3", basis="CoolProp 6.6")
    assert C.rho_hot.magnitude == pytest.approx(971.8)
    with pytest.raises(ValueError, match="already defined"):
        C.add("rho_w", 1.0, "kg/m3")
    C.override("rho_w", 958.4, "kg/m3", basis="boiling case")
    assert C.rho_w.magnitude == pytest.approx(958.4)
    # an override keeps the reviewed description unless it is replaced
    assert C.entry("rho_w").description == "Water density"
    assert C.entry("rho_w").basis == "boiling case"


def test_constants_table_uses_the_sheet_symbols_and_definitions(constants_file):
    table = Constants.load(constants_file).table("Cv_f", "d_face")
    assert [str(c.title) for c in table.columns] == [
        "Symbol", "Value / unit", "Definition / basis"]
    assert str(table.rows[0][0]) == "C_vf"      # the sheet's display symbol
    assert table.rows[1][2] == "Screen diameter; Design"


def test_format_substitutes_constants_into_prose(constants_file):
    C = Constants.load(constants_file)
    assert C.format("{d_face} screen") == "58 mm screen"
    assert C.format("{rho_w:1}") == "998 kg/m³"


# sheets


def test_a_sheet_titles_its_table_from_its_own_header_row(tmp_path):
    path = workbook(tmp_path, "r.xlsx", ["ID", "Requirement", "Assessment"],
                    [("R01", "Housing MEOP", "Met"),
                     ("R02", "Housing proof", "Met")])
    sheet = Sheet.load(path, unique="id")
    assert sheet[0]["requirement"] == "Housing MEOP"
    table = sheet.table()
    assert [str(c.title) for c in table.columns] == ["ID", "Requirement", "Assessment"]
    assert table.rows[1] == ["R02", "Housing proof", "Met"]
    assert [str(c.title) for c in sheet.table(hide=("assessment",)).columns] == [
        "ID", "Requirement"]


def test_cells_reference_constants_and_unknown_names_fail(tmp_path, constants_file):
    C = Constants.load(constants_file)
    path = workbook(tmp_path, "r.xlsx", ["ID", "Target"],
                    [("R01", "{d_face} across")])
    assert Sheet.load(path, constants=C)[0]["target"] == "58 mm across"

    bad = workbook(tmp_path, "b.xlsx", ["ID", "Target"], [("R01", "{d_nope}")])
    with pytest.raises(ValueError, match="unknown constant"):
        Sheet.load(bad, constants=C)


def test_sheet_checks_replace_hand_written_validation(tmp_path):
    path = workbook(tmp_path, "r.xlsx", ["ID", "Target"],
                    [("R01", "a"), ("R01", "")])
    with pytest.raises(ValueError, match="row 3: duplicate id"):
        Sheet.load(path, unique="id")
    with pytest.raises(ValueError, match="row 3: target is empty"):
        Sheet.load(path, required=("target",))
    with pytest.raises(ValueError, match="no column 'nope'"):
        Sheet.load(path, unique="nope")


def test_sources_merge_from_toml_and_a_workbook(tmp_path):
    (tmp_path / "sources.toml").write_text(
        '[sources.astm]\ntitle = "ASTM A240"\n', encoding="utf-8")
    path = workbook(tmp_path, "s.xlsx", ["key", "title", "url", "note"],
                    [("twp", "TWP 40 mesh", "https://example.invalid", "catalog")])
    merged = Sources.load(tmp_path / "sources.toml", path)
    assert set(merged) == {"astm", "twp"}
    assert merged["twp"].url == "https://example.invalid"
    assert merged["twp"].note == "catalog"


# display units


def test_display_unit_comments_replace_a_table_of_expected_units():
    @calculation
    def bending(P, L):
        c, I = 25 * mm, 4.0e5 * mm ** 4
        # equations
        M = P * L                       # -> kN*m
        sigma = M * c / I               # -> MPa

    result = bending(2 * ureg.kN, 100 * mm)
    assert result.M.units == ureg.kN * ureg.m
    assert result.sigma.units == ureg.MPa
    # the annotation is not rendered as a stray comment
    assert "#" not in result.source
    assert result.values["sigma"].magnitude == pytest.approx(12.5)


def test_display_units_work_in_an_inline_calculation_cell():
    doc = build(source='''from kip import *
# %% calc bending "Bending"
P = 2 * kN
L = 100 * mm
M = P * L      # -> kN*m
''')
    assert doc.value('M').units == ureg.kN * ureg.m


def test_an_impossible_display_unit_names_the_conversion():
    with pytest.raises(ExecutionError, match="cannot convert"):
        build(source='from kip import *\n# %% calc a "A"\nx = 2 * kN   # -> mm\n')


def test_a_dangling_display_unit_is_reported():
    @calculation(units={"nope": "mm"})
    def calc():
        # equations
        x = 1 * mm

    with pytest.raises(ValueError, match="does not assign"):
        calc()


# placement


def test_an_empty_content_cell_places_an_object_built_elsewhere(tmp_path):
    path = workbook(tmp_path, "r.xlsx", ["ID", "Requirement"], [("R01", "Shall")])
    doc = build(source=f'''from kip import *
reqs = Sheet.load(r"{path}").table()
# %% text intro "Intro"
"""Prose."""

# %% table reqs "Requirement assessment"
''')
    table = doc.results["reqs"].content
    assert [str(c.title) for c in table.columns] == ["ID", "Requirement"]
    assert table.rows == [["R01", "Shall"]]


def test_an_empty_cell_naming_nothing_is_a_clear_error():
    with pytest.raises(ExecutionError, match="no Table of that name"):
        build(source='from kip import *\n# %% table missing "Missing"\n')


@pytest.mark.parametrize("body", ["", "data", 'Sheet.load("values.xlsx")'])
def test_table_cells_place_sheets_without_a_table_call(tmp_path, body):
    workbook(tmp_path, "values.xlsx", ["key", "value", "unit"], [("P", 12, "kN")])
    document = build(source='from kip import *\ndata = Sheet.load("values.xlsx")\n'
                     '# %% table data "Inputs"\n' + body, path=tmp_path / "doc.py")
    assert document.results["data"].content.rows == [["P", 12, "kN"]]


def test_table_cells_place_constants_without_a_table_call(tmp_path):
    (tmp_path / "input").mkdir()
    workbook(tmp_path / "input", "constants.xlsx", ["key", "value", "unit"], [("P", 12, "kN")])
    document = build(source='from kip import *\nC = Constants.load()\n'
                     '# %% table "Inputs"\nC\n', path=tmp_path / "doc.py")
    assert document.results["inputs"].content.rows == document.value("C").table().rows


def test_relative_document_folder_is_resolved_once(tmp_path, monkeypatch):
    from pathlib import Path

    inputs = tmp_path / "project" / "input"
    inputs.mkdir(parents=True)
    workbook(inputs, "constants.xlsx", ["key", "value", "unit"], [("P", 12, "kN")])
    monkeypatch.chdir(tmp_path)
    document = build(source='from kip import *\n# %% table "Inputs"\nConstants.load()\n',
                     path=Path("project/doc.py"))
    assert document.results["inputs"].content.rows


def test_fmt_quantity_does_not_pad_or_lose_digits():
    assert fmt_quantity(650.0 * ureg.psi) == "650 psi"
    assert fmt_quantity(0.000979 * ureg.Pa) == "0.000979 Pa"
    assert fmt_quantity(997.99 * ureg.kg) == "997.99 kg"


# generated workbooks


def test_a_saved_workbook_carries_no_metadata_and_is_reproducible(tmp_path):
    import zipfile

    from kip.sheets import save_workbook

    def build(path):
        wb = Workbook()
        wb.active.title = "Inputs"
        wb.active.append(["key", "value"])
        wb.active.append(["rho", 997.99])
        return save_workbook(wb, path)

    first, second = build(tmp_path / "a.xlsx"), build(tmp_path / "b.xlsx")
    assert first.read_bytes() == second.read_bytes()

    with zipfile.ZipFile(first) as zf:
        assert {entry.date_time for entry in zf.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        core = zf.read("docProps/core.xml").decode()
        assert "openpyxl" not in core and "dcterms:modified" not in core
        assert "Openpyxl" not in zf.read("docProps/app.xml").decode()

    # still a workbook kip can read back
    assert Sheet.load(first)[0] == {"key": "rho", "value": 997.99}


def test_a_references_workbook_fills_every_source_field(tmp_path):
    path = workbook(
        tmp_path, "references.xlsx",
        ["key", "title", "author", "publisher", "year", "section", "url", "note"],
        [("armour", "Fluid Flow Through Woven Screens", "Armour and Cannon",
          "AIChE Journal", 1968, "14(3)", "https://example.invalid", "Screen model"),
         ("roark", "Roark's Formulas", "Young", "McGraw-Hill", "", "", "", "")],
    )
    refs = Sources.load(path)
    assert refs["armour"].author == "Armour and Cannon"
    assert refs["armour"].year == "1968"
    assert refs["armour"].section == "14(3)"
    # empty cells stay absent rather than becoming empty strings in the listing
    assert refs["roark"].year is None
    assert refs["roark"].full() == "Young, Roark's Formulas, McGraw-Hill"


def test_a_second_worksheet_is_read_by_name(tmp_path):
    wb = Workbook()
    wb.active.title = "Inputs"
    wb.active.append(["key", "title"])
    wb.active.append(["a", "A"])
    docs = wb.create_sheet("Documents")
    for row in (["Document", "Organization", "Number"],
                ["Multi-pass filter evaluation", "ISO", "16889:2022"]):
        docs.append(row)
    path = tmp_path / "references.xlsx"
    wb.save(path)

    sheet = Sheet.load(path, "Documents", unique="number")
    assert [str(c.title) for c in sheet.table().columns] == [
        "Document", "Organization", "Number"]
    assert sheet[0]["organization"] == "ISO"
