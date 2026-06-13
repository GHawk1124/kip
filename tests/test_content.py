"""Figures, tables, drawings and sources: data model and rendering."""
import pytest
import typst

from kip.content import (
    Column, Drawing, Figure, Source, Sources, Table, strip_units,
)
from kip.doc import build
from kip.render import emit
from kip.render.emitter import _figure, _mark, _short_url
from kip.render.layout import Layout
from kip.units import ureg


# --- figures ---------------------------------------------------------------

def test_strip_units_converts_before_stripping():
    vals = [1 * ureg.m, 500 * ureg.mm]
    assert strip_units(vals, "mm") == [1000.0, 500.0]


def test_figure_infers_axis_units_from_quantities():
    fig = Figure(xlabel="Thickness", ylabel="Stress")
    fig.line([8 * ureg.mm, 16 * ureg.mm], [100 * ureg.MPa, 50 * ureg.MPa])
    assert fig.xlabel == "Thickness (mm)"
    assert fig.ylabel == "Stress (MPa)"
    assert fig.series[0].x == [8.0, 16.0]


def test_figure_of_plots_a_sympy_expression():
    import sympy as sp

    x = sp.Symbol("x")
    fig = Figure().of(x ** 2, x, 0, 3, n=4)
    assert fig.series[0].y == pytest.approx([0.0, 1.0, 4.0, 9.0])


def test_unknown_mark_raises_a_python_error():
    """Better than a Typst assertion failing deep inside lilaq."""
    with pytest.raises(ValueError, match="unknown plot mark"):
        _mark("hexagram")


@pytest.mark.parametrize("alias,expected", [
    ("square", '"s"'), ("circle", '"o"'), ("triangle", '"^"'),
    ("diamond", '"d"'), (None, "none"),
])
def test_mark_aliases(alias, expected):
    assert _mark(alias) == expected


def test_figure_renders_valid_typst():
    fig = Figure(xlabel="x", ylabel="y", width=80, height=50)
    fig.line([0, 1, 2], [0, 1, 4], label="a", mark="o")
    fig.scatter([0, 1], [1, 2], label="b", mark="square")
    src = ('#import "@preview/lilaq:0.5.0" as lq\n'
           "#set page(width: 120mm, height: auto)\n" + _figure(fig))
    typst.compile({"main.typ": src.encode()}, timestamp=0)


# --- tables ----------------------------------------------------------------

def make_table(**kw):
    return Table(
        columns=[Column("a", "A", align="left"),
                 Column("b", "Load", unit="kN", precision=1)],
        rows=[("x", 1000 * ureg.N), ("y", 2500 * ureg.N)],
        **kw,
    )


def test_table_converts_units_per_column():
    rows, hidden = make_table().display_rows()
    assert rows == [["x", "1.0"], ["y", "2.5"]]
    assert hidden == 0


def test_table_header_carries_units():
    assert make_table().headers == ["A", "Load (kN)"]


def test_max_rows_truncates_the_render_only():
    tbl = make_table(max_rows=1)
    rows, hidden = tbl.display_rows()
    assert len(rows) == 1 and hidden == 1
    assert len(tbl.rows) == 2, "the data itself must not be truncated"


def test_xlsx_export_keeps_every_row_and_writes_numbers(tmp_path):
    from openpyxl import load_workbook

    tbl = make_table(max_rows=1, xlsx="out.xlsx")
    path = tbl.to_xlsx(tmp_path / "out.xlsx")
    ws = load_workbook(path).active

    assert ws.max_row == 3, "header + both rows, despite max_rows=1"
    assert [c.value for c in ws[1]] == ["A", "Load (kN)"]
    # numbers must be numbers, not strings, or the spreadsheet is useless
    assert ws.cell(row=2, column=2).value == pytest.approx(1.0)
    assert ws.cell(row=3, column=2).value == pytest.approx(2.5)


# --- sources ---------------------------------------------------------------

def test_source_formats_short_and_full():
    s = Source(title="ASME BTH-1", section="3-3.2", publisher="ASME", year=2023)
    assert s.short() == "ASME BTH-1 3-3.2"
    assert "ASME" in s.full() and "2023" in s.full()


def test_sources_accepts_dicts_and_objects():
    srcs = Sources({"a": {"title": "T"}}, b=Source(title="U"))
    assert isinstance(srcs["a"], Source) and srcs["b"].title == "U"


@pytest.mark.parametrize("url,expected", [
    ("https://example.com/a", "example.com/a"),
    ("http://example.com/", "example.com"),
    (None, ""),
])
def test_short_url(url, expected):
    assert _short_url(url) == expected


def test_short_url_elides_long_tails():
    out = _short_url("https://x.com/" + "a" * 200)
    assert len(out) <= 48 and out.endswith("…") and out.startswith("x.com/")


# --- end to end ------------------------------------------------------------

DOC = r'''from kip import *

# %% kip.plot id=fig label="A plot"
fig = Figure(xlabel="x", ylabel="y", width=80, height=45)
fig.line([0, 1, 2], [0, 1, 4], label="curve", mark="o")

# %% kip.table id=tbl label="A table"
tbl = Table(columns=["Name", "Value"], rows=[("a", 1.0), ("b", 2.0)])

# %% kip.draw id=dwg label="A sketch"
dwg = Drawing(body="circle((0,0), radius: 1)\n  line((-1,0), (1,0))")

# %% kip.sources id=refs
refs = Sources(spec=Source(title="A Standard", url="https://example.com/spec"))

# %% kip.text id=note
"""See @src:spec and @blk:fig."""
'''


def test_all_content_kinds_build_and_render(tmp_path):
    doc = build(source=DOC, path="doc.py")
    assert not doc.errors
    assert all(r.ok for r in doc.results.values()), [
        (r.block_id, r.error) for r in doc.results.values() if r.failed
    ]
    assert isinstance(doc.results["fig"].content, Figure)
    assert isinstance(doc.results["tbl"].content, Table)
    assert isinstance(doc.results["dwg"].content, Drawing)
    assert isinstance(doc.results["refs"].content, Sources)

    from kip.render.pdf import compile_pdf

    pdf = compile_pdf(emit(doc, Layout()))
    assert len(pdf) > 1000


def test_content_block_without_the_right_object_fails_clearly():
    src = '''from kip import *

# %% kip.plot id=fig
x = 42
'''
    doc = build(source=src, path="doc.py", strict=False)
    assert doc.results["fig"].failed
    assert "must bind a Figure" in doc.results["fig"].error


def test_content_kinds_are_not_linted_as_handcalcs():
    """Method calls are fine in a plot block; they are not rendered by handcalcs."""
    src = '''from kip import *

# %% kip.plot id=fig
fig = Figure()
fig.line([0, 1], [0, 1])
'''
    doc = build(source=src, path="doc.py", strict=False)
    assert not doc.errors
