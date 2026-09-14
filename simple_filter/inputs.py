"""Read the three editable Excel inputs, without cached defaults."""
import math
from openpyxl import load_workbook
from pathlib import Path
ROOT=Path(__file__).resolve().parent

def read_input(name, required):
    path=ROOT/'input'/f'{name}.xlsx'
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        values=list(book['Inputs'].iter_rows(values_only=True))
        headers=values[0]
        if len(set(headers))!=len(headers) or not set(required).issubset(headers):
            raise ValueError(f'{path.name}: missing or duplicate columns')
        rows=[]
        for row in values[1:]:
            if not any(v is not None for v in row): continue
            if any(isinstance(v,str) and v.startswith('=') for v in row):
                raise ValueError(f'{path.name}: enter values, not formulas')
            rows.append({k:v if v is not None else '' for k,v in zip(headers,row)})
        return rows
    finally: book.close()

ROWS=read_input('constants',['key','value','unit','description','basis'])
C={}
for row in ROWS:
    key=row['key']
    if key in C or not isinstance(row['value'],(int,float)) or not math.isfinite(row['value']):
        raise ValueError(f'Invalid or duplicate constant: {key}')
    C[key]=float(row['value'])
EXPECTED_UNITS={'temperature_F': 'degF', 'rho_w': 'kg/m3', 'mu_w': 'Pa s', 'mdot_n': 'lb/s', 'mdot_s': 'lb/s', 'p_meop': 'psi', 'p_proof': 'psi', 'p_burst': 'psi', 'dp_clean_limit': 'psi', 'dp_loaded_limit': 'psi', 'dirt_mass': 'g', 'mass_limit': 'lb', 'rating': 'um', 'efficiency': 'percent', 'efficiency_size': 'um', 'active_d': 'mm', 'frame_d': 'mm', 'frame_raw_t': 'mm', 'frame_assembled_t': 'mm', 'coarse_t': 'mm', 'fine_t': 'mm', 'body_d': 'mm', 'pocket_d': 'mm', 'outer_shoulder_z': 'mm', 'inner_shoulder_z': 'mm', 'stub_start_z': 'mm', 'half_length': 'mm', 'stub_od': 'mm', 'stub_id': 'mm', 'Cv_f': '1/m', 'Ci_f': '1', 'Cv_c': '1/m', 'Ci_c': '1', 'K_body': '1', 'alpha_c': 'm/kg', 'rho_ti': 'kg/m3', 'rho_ss': 'kg/m3', 'Sy': 'MPa', 'Su': 'MPa', 'coarse_areal_mass': 'kg/m2', 'fine_areal_mass': 'kg/m2', 'growth_factor': '1'}
for row in ROWS:
    if row["key"] in EXPECTED_UNITS and row["unit"] != EXPECTED_UNITS[row["key"]]:
        raise ValueError(f"Use {EXPECTED_UNITS[row['key']]} for {row['key']}")
REQS=read_input('requirements',['id','requirement','constant_key','qualifier_key','target_text','assessment'])
SUPPLIERS=read_input('mesh_suppliers',['supplier_item','catalog_basis','assessment','source_key','source_url'])
if len({r['id'] for r in REQS})!=len(REQS): raise ValueError('Duplicate requirement ID')
for row in REQS:
    for field in ['constant_key','qualifier_key']:
        if row[field] and row[field] not in C: raise ValueError(f'Unknown constant: {row[field]}')
    if not row['constant_key'] and not row['target_text']: raise ValueError(f'Missing target: {row["id"]}')
if len({r['source_key'] for r in SUPPLIERS})!=len(SUPPLIERS): raise ValueError('Duplicate supplier source key')
LB_KG=0.45359237
PSI_PA=6894.757293168
